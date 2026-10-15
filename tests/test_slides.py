import io
import json

from PIL import Image
from pptx import Presentation

from bot.files.blocks import OutFile
from bot.files.slides import THEMES, build_pptx, build_preview, parse_deck, render_decks


def png(w: int = 64, h: int = 32) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (w, h), "#c33").save(out, "PNG")
    return out.getvalue()


SPEC = {
    "title": "Кофейня «Зерно»",
    "theme": "dark",
    "slides": [
        {"layout": "title", "title": "Кофейня «Зерно»", "subtitle": "План запуска"},
        {"layout": "section", "title": "Рынок"},
        {
            "layout": "bullets",
            "title": "Почему сейчас",
            "bullets": ["Спрос растёт", "Аренда упала"],
            "notes": "Цифры из отчёта",
        },
        {
            "layout": "two_columns",
            "title": "Сравнение",
            "left": {"heading": "Мы", "bullets": ["Свой обжар"]},
            "right": {"heading": "Они", "bullets": ["Сеть"]},
        },
        {"layout": "stat", "title": "Окупаемость", "value": "14 мес", "caption": "при 120 чашках в день"},
        {"layout": "quote", "quote": "Лучший кофе — свежий", "author": "Бариста"},
        {"layout": "image", "title": "Интерьер", "image": "hall.png", "bullets": ["Светлый зал"]},
        {"layout": "mystery", "title": "Неизвестный вид <script>", "bullets": ["станет списком"]},
        {"layout": "bullets"},  # empty: dropped
    ],
}


def test_deck_builds_a_16x9_pptx_with_every_slide_and_notes() -> None:
    deck = parse_deck(json.dumps(SPEC, ensure_ascii=False))
    assert len(deck.slides) == 8 and deck.theme is THEMES["dark"]
    assert deck.slides[-1].layout == "bullets"  # unknown layout falls back

    prs = Presentation(io.BytesIO(build_pptx(deck, {"hall.png": png()})))
    assert len(prs.slides) == 8
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    texts = [" ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame) for slide in prs.slides]
    assert "План запуска" in texts[0] and "Спрос растёт" in texts[2] and "14 мес" in texts[4]
    assert prs.slides[2].notes_slide.notes_text_frame.text == "Цифры из отчёта"
    assert any(s.shape_type == 13 for s in prs.slides[6].shapes)  # 13 = picture


def test_preview_has_one_frame_per_slide_and_escapes_text() -> None:
    deck = parse_deck(json.dumps(SPEC, ensure_ascii=False))
    page = build_preview(deck, {"hall.png": png()})
    assert page.count('class="cell"') == 8
    assert "aspect-ratio:16/9" in page and "data:image/png;base64," in page
    assert "Неизвестный вид &lt;script&gt;" in page


def test_image_slide_without_the_picture_becomes_a_list() -> None:
    deck = parse_deck(
        json.dumps({"slides": [{"layout": "image", "title": "Т", "image": "нет.png", "bullets": ["п"]}]})
    )
    prs = Presentation(io.BytesIO(build_pptx(deck, {})))
    assert not any(s.shape_type == 13 for s in prs.slides[0].shapes)


def test_render_decks_replaces_the_block_and_keeps_a_broken_one_as_json() -> None:
    good = OutFile("plan.pptx", json.dumps(SPEC, ensure_ascii=False).encode())
    bad = OutFile("bad.pptx", b"{not json")
    other = OutFile("notes.md", b"# hi")
    names = [f.name for f in render_decks([good, bad, other], {"hall.png": png()})]
    assert names == ["plan.pptx", "plan (превью).html", "bad.json", "notes.md"]


def test_fenced_json_inside_the_block_is_accepted() -> None:
    deck = parse_deck("```json\n" + json.dumps({"slides": [{"title": "А", "bullets": ["б"]}]}) + "\n```")
    assert deck.slides[0].data["bullets"] == ["б"]


def test_fit_takes_the_largest_size_that_fits() -> None:
    from bot.files.slides import BULLET_SIZES, _fit

    assert _fit(["Коротко"], 11.7, 4.9, BULLET_SIZES) == BULLET_SIZES[0]  # capped, not huge
    five = ["Кофе навынос и локальные обжарщики растут быстрее общепита в целом"] * 5
    middle = _fit(five, 11.7, 4.9, BULLET_SIZES, indent=1.6)
    assert BULLET_SIZES[-1] < middle < BULLET_SIZES[0]  # long bullets shrink, but stay readable
    assert _fit(["слово " * 400], 5.0, 1.0, BULLET_SIZES) == BULLET_SIZES[-1]  # never below the floor
