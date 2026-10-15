from __future__ import annotations

import io
import zipfile

import pytest

from bot.files.extract import MAX_TEXT_CHARS, UnsupportedFile, extract_text


def tiny_pdf(text: str) -> bytes:
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 200] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 18 Tf 40 120 Td ({text}) Tj ET".encode()
    objs[3] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


DOCX_DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>
<w:p><w:r><w:t>Hello</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>World</w:t></w:r>
<w:r><w:br/></w:r><w:r><w:t>Second line</w:t></w:r></w:p>
<w:tbl>
<w:tr><w:tc><w:p><w:r><w:t>A1</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>B1</w:t></w:r></w:p></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:t>A2</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>B2</w:t></w:r></w:p></w:tc></w:tr>
</w:tbl>
</w:body>
</w:document>
"""


def make_docx(document_xml: str = DOCX_DOCUMENT_XML) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="x"></Types>',
        )
        zf.writestr("word/document.xml", document_xml)
    return buf.getvalue()


WORKBOOK_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets>
</workbook>
"""

WORKBOOK_RELS_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>
"""

SHARED_STRINGS_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1">
<si><t>Shared</t></si>
</sst>
"""

SHEET1_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData>
<row r="1"><c r="A1" t="s"><v>0</v></c><c r="C1" t="inlineStr"><is><t>Inline</t></is></c></row>
<row r="2"><c r="A2" t="b"><v>1</v></c><c r="B2"><v>42</v></c></row>
</sheetData>
</worksheet>
"""


def make_xlsx(
    *,
    workbook_xml: str = WORKBOOK_XML,
    rels_xml: str | None = WORKBOOK_RELS_XML,
    shared_xml: str | None = SHARED_STRINGS_XML,
    sheet1_xml: str = SHEET1_XML,
) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml", workbook_xml)
        if rels_xml is not None:
            zf.writestr("xl/_rels/workbook.xml.rels", rels_xml)
        if shared_xml is not None:
            zf.writestr("xl/sharedStrings.xml", shared_xml)
        zf.writestr("xl/worksheets/sheet1.xml", sheet1_xml)
    return buf.getvalue()


# --- PDF ----------------------------------------------------------------------


def test_pdf_extracts_text():
    result = extract_text("doc.pdf", tiny_pdf("Hello PDF"))
    assert result.kind == "pdf"
    assert "Hello PDF" in result.text
    assert "--- страница 1 ---" in result.text
    assert not result.truncated


def test_pdf_sniffed_without_extension():
    result = extract_text("doc.bin", tiny_pdf("Sniffed"))
    assert result.kind == "pdf"
    assert "Sniffed" in result.text


def test_pdf_blank_text_layer_raises():
    with pytest.raises(UnsupportedFile, match="текстового слоя"):
        extract_text("scan.pdf", tiny_pdf(""))


def test_pdf_corrupted_raises():
    with pytest.raises(UnsupportedFile, match="прочитать PDF"):
        extract_text("bad.pdf", b"%PDF-1.4\nnot really a pdf")


# --- DOCX ----------------------------------------------------------------------


def test_docx_paragraphs_tabs_breaks_and_table():
    result = extract_text("doc.docx", make_docx())
    assert result.kind == "docx"
    lines = result.text.split("\n")
    assert lines[0] == "Hello\tWorld"
    assert lines[1] == "Second line"
    assert "A1 | B1" in result.text
    assert "A2 | B2" in result.text


def test_docx_corrupted_zip_raises():
    with pytest.raises(UnsupportedFile, match="DOCX"):
        extract_text("bad.docx", b"not a zip at all")


def test_docx_missing_document_xml_raises():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("word/other.xml", "<x/>")
    with pytest.raises(UnsupportedFile, match="DOCX"):
        extract_text("bad.docx", buf.getvalue())


# --- XLSX ----------------------------------------------------------------------


def test_xlsx_shared_inline_and_gap_column():
    result = extract_text("book.xlsx", make_xlsx())
    assert result.kind == "xlsx"
    assert "## Лист: Data" in result.text
    lines = result.text.split("\n")
    assert "Shared,,Inline" in lines
    assert "TRUE,42" in lines


def test_xlsx_corrupted_zip_raises():
    with pytest.raises(UnsupportedFile, match="XLSX"):
        extract_text("bad.xlsx", b"not a zip at all")


# --- text files ------------------------------------------------------------------


def test_text_utf8():
    result = extract_text("notes.txt", "Привет мир".encode())
    assert result.kind == "text"
    assert result.text == "Привет мир"


def test_text_utf8_bom():
    result = extract_text("notes.txt", "Привет".encode("utf-8-sig"))
    assert result.text == "Привет"


def test_text_cp1251():
    result = extract_text("notes.txt", "Привет".encode("cp1251"))
    assert result.text == "Привет"


def test_text_unknown_extension_py_csv_json():
    for name in ("script.py", "data.csv", "data.json", "notes.md", "app.log", "README"):
        result = extract_text(name, b"plain content")
        assert result.kind == "text"
        assert result.text == "plain content"


def test_binary_blob_with_nul_bytes_raises():
    with pytest.raises(UnsupportedFile):
        extract_text("data.bin", b"\x00\x01\x02\x03" * 100)


def test_undecodable_bytes_raise():
    with pytest.raises(UnsupportedFile):
        extract_text("weird.txt", b"\xff\xfe\x00\xff" * 10)


# --- known unsupported binaries --------------------------------------------------


@pytest.mark.parametrize("ext", [".doc", ".xls", ".ppt", ".rar", ".7z", ".exe", ".mp3", ".mp4"])
def test_known_binary_extensions_raise_with_hint(ext):
    with pytest.raises(UnsupportedFile, match="PDF, DOCX, XLSX, PPTX"):
        extract_text(f"file{ext}", b"whatever bytes")


def test_pptx_slides_in_order_with_notes():
    from pptx import Presentation

    prs = Presentation()
    for n in range(1, 12):  # 11 slides: slide10 must come after slide9, not after slide1
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Заголовок {n}"
        slide.placeholders[1].text = f"Пункт слайда {n}"
    prs.slides[1].notes_slide.notes_text_frame.text = "Сказать про бюджет"
    out = io.BytesIO()
    prs.save(out)
    result = extract_text("deck.pptx", out.getvalue())
    assert result.kind == "pptx"
    order = [line for line in result.text.splitlines() if line.startswith("Заголовок")]
    assert order == [f"Заголовок {n}" for n in range(1, 12)]
    assert "Пункт слайда 7" in result.text
    assert "Заметки: Сказать про бюджет" in result.text


def test_broken_pptx_is_a_clear_refusal():
    with pytest.raises(UnsupportedFile, match="PPTX"):
        extract_text("deck.pptx", b"not a zip at all")


# --- normalization ----------------------------------------------------------------


def test_crlf_normalized_to_lf():
    result = extract_text("notes.txt", b"line1\r\nline2\r\n")
    assert result.text == "line1\nline2"


def test_truncation_over_max_chars():
    data = ("a" * (MAX_TEXT_CHARS + 500)).encode()
    result = extract_text("big.txt", data)
    assert result.truncated is True
    assert len(result.text) == MAX_TEXT_CHARS


# --- zip archives ------------------------------------------------------------------


def _zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def test_zip_of_a_project_reads_every_text_member_under_its_path():
    data = _zip(
        {
            "app/main.py": "print('привет')".encode(),
            "app/README.md": "# Проект".encode(),
            "docs/spec.pdf": tiny_pdf("Spec inside"),
            "app/logo.png": b"\x89PNG\r\n\x1a\n\x00\x00",
            "node_modules/lib/index.js": b"ignored()",
            "app/__pycache__/main.cpython-314.pyc": b"\x00\x01",
        }
    )
    result = extract_text("project.zip", data)
    assert result.kind == "zip"
    assert "=== app/main.py ===\nprint('привет')" in result.text
    assert "=== app/README.md ===\n# Проект" in result.text
    assert "Spec inside" in result.text
    assert "app/logo.png" in result.text.split("не прочитаны")[1]  # listed, not read
    assert "node_modules" not in result.text and "__pycache__" not in result.text


def test_zip_without_readable_members_or_broken_raises():
    with pytest.raises(UnsupportedFile, match="умею читать"):
        extract_text("pics.zip", _zip({"a.png": b"\x89PNG\x00"}))
    with pytest.raises(UnsupportedFile, match="ZIP"):
        extract_text("bad.zip", b"not a zip")
